import os
import re
import shutil
import pdfplumber
import openpyxl
from openpyxl.styles import Font, PatternFill
from pypdf import PdfReader, PdfWriter

def extract_page_data(pdf_path, folder_name):
    """Scans a PDF and returns a list of dictionaries with extracted data for each page."""
    page_data = []
    
    file_name_lower = os.path.basename(pdf_path).lower()
    
    is_folder_capella = "capella" in folder_name.lower()
    is_folder_flavor_west = "flavor west" in folder_name.lower()
    is_folder_hedessent = "hedessent" in folder_name.lower() or "flavourart" in folder_name.lower() or "flavour art" in folder_name.lower()
    is_folder_flavorah = "flavorah" in folder_name.lower()
    is_folder_lorann = bool(re.search(r'lorann|(?<![a-z])la(?![a-z])', folder_name.lower()))
    is_folder_purilum = "purilum" in folder_name.lower()
    is_folder_sobucky = bool(re.search(r'sobucky|super\s*aromas|(?<![a-z])ssa(?![a-z])', folder_name.lower()))
    is_folder_tfa = bool(re.search(r'perfumer|apprentice|(?<![a-z])tfa(?![a-z])|(?<![a-z])tpa(?![a-z])', folder_name.lower()))
    is_folder_wf = bool(re.search(r'wonder\s*flavou?rs|(?<![a-z])wf(?![a-z])', folder_name.lower()))
    
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                
                is_capella = is_folder_capella or "capella" in file_name_lower or bool(re.search(r'capella', text, re.IGNORECASE))
                is_flavor_west = is_folder_flavor_west or "flavor west" in file_name_lower or bool(re.search(r'flavor\s*west', text, re.IGNORECASE))
                is_hedessent = is_folder_hedessent or "hedessent" in file_name_lower or "flavourart" in folder_name.lower() or bool(re.search(r'hedessent|flavour\s*art|flavourart', text, re.IGNORECASE))
                is_flavorah = is_folder_flavorah or "flavorah" in file_name_lower or bool(re.search(r'flavorah', text, re.IGNORECASE))
                is_lorann = is_folder_lorann or bool(re.search(r'lorann|(?<![a-z])la(?![a-z])', file_name_lower)) or bool(re.search(r'lorann', text, re.IGNORECASE))
                is_purilum = is_folder_purilum or "purilum" in file_name_lower or bool(re.search(r'purilum', text, re.IGNORECASE))
                is_sobucky = (
                    is_folder_sobucky or 
                    bool(re.search(r'sobucky|super\s*aromas|(?<![a-z])ssa(?![a-z])', file_name_lower)) or 
                    bool(re.search(r'sobucky|super\s*aromas|\bssa\b', text, re.IGNORECASE))
                )
                is_tfa = (
                    is_folder_tfa or
                    bool(re.search(r'perfumer|apprentice|(?<![a-z])tfa(?![a-z])|(?<![a-z])tpa(?![a-z])', file_name_lower)) or
                    bool(re.search(r'perfumer\'?s\s*apprentice|flavor\s*apprentice', text, re.IGNORECASE))
                )
                is_wf = (
                    is_folder_wf or
                    bool(re.search(r'wonder\s*flavou?rs|(?<![a-z])wf(?![a-z])', file_name_lower)) or
                    bool(re.search(r'wonder\s*flavou?rs', text, re.IGNORECASE))
                )
                
                # --- UNIVERSAL EXTRACTIONS ---
                mfr_code_match = re.search(r'(?:Manufacture|Mfr|Mfg)\s*Code\s*[:\.]?\s*([A-Za-z0-9-]+)', text, re.IGNORECASE)
                mfr_code = mfr_code_match.group(1).strip() if mfr_code_match else "N/A"

                mfr_date_match = re.search(r'(?:Manufacture|Mfr|Mfg)\s*Date\s*[:\.]?\s*([A-Za-z0-9/\-\s]+?)(?=\n|$)', text, re.IGNORECASE)
                mfr_date = mfr_date_match.group(1).strip() if mfr_date_match else "N/A"

                bb_date_match = re.search(r'(?:Best\s*By|Exp(?:iration)?\s*Date|Best\s*Before)\s*[:\.]?\s*([A-Za-z0-9/\-\s]+?)(?=\n|$)', text, re.IGNORECASE)
                bb_date = bb_date_match.group(1).strip() if bb_date_match else "N/A"

                flavor_match = re.search(r'\b(?:Product\s*Description|Product|Flavor|Item|Description)\b\s*[:\.=\-]?\s*([A-Za-z0-9\-\s\(\)\'\&\+,\.%#]{2,60}?)(?=\n|  |$|\s+Ingredients|\s+approved|\s+Item|\s+Lot)', text, re.IGNORECASE)
                raw_flavor = flavor_match.group(1).strip() if flavor_match else "Unknown_Flavor"
                universal_clean_flavor = re.sub(r'[\\/*?:"<>|]', "", raw_flavor).strip().replace(" ", "_")

                if is_capella:
                    lot_match = re.search(r'(?:Lot|Batch|Lab)\s*(?:Number|No|#)?\s*[:\.]?\s*([A-Za-z0-9-]+)', text, re.IGNORECASE)
                    date_match = re.search(r'(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})', text)
                    lot = lot_match.group(1).strip() if lot_match else None
                    raw_date = date_match.group(1).strip() if date_match else "Unknown_Date"
                    clean_date = raw_date.replace("/", "-")
                    
                    page_data.append({
                        'is_capella': True, 'is_flavor_west': False, 'is_hedessent': False, 'is_flavorah': False, 'is_lorann': False, 'is_purilum': False, 'is_sobucky': False, 'is_tfa': False, 'is_wf': False,
                        'lot': lot, 'flavor': universal_clean_flavor, 'date': clean_date,
                        'mfr_code': mfr_code, 'mfr_date': mfr_date, 'bb_date': bb_date
                    })

                elif is_flavor_west:
                    item_match = re.search(r'FW-\s*([A-Za-z0-9]+)', text, re.IGNORECASE)
                    date_match = re.search(r'MANUFACTURE DATE[\s:|]*(\d{1,2})\s*[-/]\s*(\d{1,2})\s*[-/]\s*(\d{2,4})', text, re.IGNORECASE)

                    if item_match and date_match:
                        item_suffix = item_match.group(1).strip()
                        month, day, year = date_match.groups()
                        constructed_lot = f"{item_suffix}{month.zfill(2)}{day.zfill(2)}{year[-2:]}"
                        final_date = f"{month.zfill(2)}-{day.zfill(2)}-{year}"
                    else:
                        constructed_lot = None
                        final_date = "Unknown_Date"

                    fw_flavor_match = re.search(r'PRODUCT[\s:|]*([^\n]+)', text, re.IGNORECASE)
                    if fw_flavor_match:
                        raw_fw_flavor = fw_flavor_match.group(1).strip()
                        clean_fw_flavor = re.sub(r'(?i)^(?:natural\b|artifici[ea]+l\b|art\.?\b|n&a\b|and\b|&|/|-|\s)+', '', raw_fw_flavor).strip()
                        clean_fw_flavor = re.sub(r'(?i)\s+Flavor\.?$', '', clean_fw_flavor).strip()
                        fw_clean_flavor = re.sub(r'[\\/*?:"<>|]', "", clean_fw_flavor).strip().replace(" ", "_")
                    else:
                        fw_clean_flavor = universal_clean_flavor 

                    page_data.append({
                        'is_capella': False, 'is_flavor_west': True, 'is_hedessent': False, 'is_flavorah': False, 'is_lorann': False, 'is_purilum': False, 'is_sobucky': False, 'is_tfa': False, 'is_wf': False,
                        'lot': constructed_lot, 'flavor': fw_clean_flavor, 'date': final_date,
                        'mfr_code': mfr_code, 'mfr_date': mfr_date, 'bb_date': bb_date
                    })

                elif is_hedessent:
                    is_flavourart_layout = bool(re.search(r'Lotto\s*/\s*Batch|Aroma\s*/\s*Flavour', text, re.IGNORECASE))
                    if is_flavourart_layout:
                        fa_lot_match = re.search(r'Lotto\s*/\s*Batch[\s:|]*([A-Za-z0-9-]+)', text, re.IGNORECASE)
                        lot = fa_lot_match.group(1).strip() if fa_lot_match else None
                        fa_prod_match = re.search(r'Aroma\s*/\s*Flavour[\s:|]*(?:([A-Za-z0-9-]+)\s*-\s*)?(?:([^/\n]+)/\s*)?([^\n]+)', text, re.IGNORECASE)
                        if fa_prod_match:
                            hed_mfr_code = fa_prod_match.group(1).strip() if fa_prod_match.group(1) else mfr_code
                            raw_flavor = fa_prod_match.group(3).strip() 
                            clean_hed_flavor = re.sub(r'(?i)\s+Flavour\.?$', '', raw_flavor).strip()
                            clean_hed_flavor = re.sub(r'(?i)\bImitation\b', '', clean_hed_flavor).strip()
                            hed_clean_flavor = re.sub(r'[\\/*?:"<>|]', "", clean_hed_flavor).strip().replace(" ", "_")
                        else:
                            hed_mfr_code = mfr_code
                            hed_clean_flavor = universal_clean_flavor
                            
                        fa_mfg_match = re.search(r'Date of manufacture[\s:]+(?:[^/]+/\s*)?([^\n]+)', text, re.IGNORECASE)
                        hed_mfr_date = fa_mfg_match.group(1).strip() if fa_mfg_match else mfr_date
                        fa_retest_match = re.search(r'Retest date[\s:]+(?:[^/]+/\s*)?([^\n]+)', text, re.IGNORECASE)
                        hed_bb_date = fa_retest_match.group(1).strip() if fa_retest_match else bb_date
                        clean_date = hed_mfr_date.replace("/", "-").replace(" ", "-") if hed_mfr_date != "N/A" else "Unknown_Date"
                    else:
                        hed_lot_match = re.search(r'LOT[\s\nA-Z0-9]*?([0-9]{6,}-[0-9]+)', text, re.IGNORECASE)
                        if not hed_lot_match:
                            hed_lot_match = re.search(r'(?:Lot|Batch)\s*(?:Number|No|#)?\s*[:\.]?\s*([A-Za-z0-9-]+)', text, re.IGNORECASE)
                        lot = hed_lot_match.group(1).strip() if hed_lot_match else None

                        hed_code_match = re.search(r'Code[\s\n]+(?:LOT[\s\n]+)?([A-Z]{2,4}[0-9]+)', text, re.IGNORECASE)
                        hed_mfr_code = hed_code_match.group(1).strip() if hed_code_match else mfr_code
                        hed_mfg_date_match = re.search(r'DATE OF MANUFACTURE[\s:]+([^\n]+)', text, re.IGNORECASE)
                        hed_mfr_date = hed_mfg_date_match.group(1).strip() if hed_mfg_date_match else mfr_date
                        hed_bb_date_match = re.search(r'RETEST DATE[\s:]+([^\n]+)', text, re.IGNORECASE)
                        hed_bb_date = hed_bb_date_match.group(1).strip() if hed_bb_date_match else bb_date
                        clean_date = hed_mfr_date.replace("/", "-").replace(" ", "-") if hed_mfr_date != "N/A" else "Unknown_Date"

                        hed_flavor_match = re.search(r'PRODUCT[\s\n]+([^\|\n]+)', text, re.IGNORECASE)
                        if hed_flavor_match:
                            raw_hed_flavor = hed_flavor_match.group(1).strip()
                            clean_hed_flavor = re.sub(r'(?i)\s+Flavour\.?$', '', raw_hed_flavor).strip()
                            clean_hed_flavor = re.sub(r'(?i)\bImitation\b', '', clean_hed_flavor).strip()
                            hed_clean_flavor = re.sub(r'[\\/*?:"<>|]', "", clean_hed_flavor).strip().replace(" ", "_")
                        else:
                            hed_clean_flavor = universal_clean_flavor

                    page_data.append({
                        'is_capella': False, 'is_flavor_west': False, 'is_hedessent': True, 'is_flavorah': False, 'is_lorann': False, 'is_purilum': False, 'is_sobucky': False, 'is_tfa': False, 'is_wf': False,
                        'lot': lot, 'flavor': hed_clean_flavor, 'date': clean_date,
                        'mfr_code': hed_mfr_code, 'mfr_date': hed_mfr_date, 'bb_date': hed_bb_date
                    })

                elif is_flavorah:
                    flv_lot_code_match = re.search(r'Lot\s*/\s*Batch\s*Number:[\s\n\|]*Product\s*ID\s*Number:[\s\n\|]*([A-Za-z0-9-]+)[\s\n\|]+([A-Za-z0-9-]+)', text, re.IGNORECASE)
                    if flv_lot_code_match:
                        lot = flv_lot_code_match.group(1).strip()
                        flv_mfr_code = flv_lot_code_match.group(2).strip()
                    else:
                        lot_match_solo = re.search(r'Lot\s*/\s*Batch\s*Number:[\s\n\|]*(?!Product\b)([A-Za-z0-9-]+)', text, re.IGNORECASE)
                        lot = lot_match_solo.group(1).strip() if lot_match_solo else None
                        code_match_solo = re.search(r'Product\s*ID\s*Number:[\s\n\|]*([A-Za-z0-9-]+)', text, re.IGNORECASE)
                        flv_mfr_code = code_match_solo.group(1).strip() if code_match_solo else mfr_code

                    flv_flavor_match = re.search(r'Flavor Name:[\s\n\|]*([^\n]+)', text, re.IGNORECASE)
                    if flv_flavor_match:
                        raw_flv_flavor = flv_flavor_match.group(1).strip()
                        clean_flv_flavor = re.sub(r'(?i)\s+(?:Type\s+)?Flavor\.?$', '', raw_flv_flavor).strip()
                        flv_clean_flavor = re.sub(r'[\\/*?:"<>|]', "", clean_flv_flavor).strip().replace(" ", "_")
                    else:
                        flv_clean_flavor = universal_clean_flavor
                        
                    flv_mfg_match = re.search(r'Date of Manufacture:[\s\n\|]*([0-9/-]+)', text, re.IGNORECASE)
                    flv_mfr_date = flv_mfg_match.group(1).strip() if flv_mfg_match else mfr_date
                    flv_bb_match = re.search(r'Best By Date:[\s\n\|]*([0-9/-]+)', text, re.IGNORECASE)
                    flv_bb_date = flv_bb_match.group(1).strip() if flv_bb_match else bb_date
                    clean_date = flv_mfr_date.replace("/", "-") if flv_mfr_date != "N/A" else "Unknown_Date"

                    page_data.append({
                        'is_capella': False, 'is_flavor_west': False, 'is_hedessent': False, 'is_flavorah': True, 'is_lorann': False, 'is_purilum': False, 'is_sobucky': False, 'is_tfa': False, 'is_wf': False,
                        'lot': lot, 'flavor': flv_clean_flavor, 'date': clean_date,
                        'mfr_code': flv_mfr_code, 'mfr_date': flv_mfr_date, 'bb_date': flv_bb_date
                    })

                elif is_lorann:
                    lor_lot_match = re.search(r'Lot\s*Code[\s\n\|:]*([A-Za-z0-9-]+)', text, re.IGNORECASE)
                    if not lor_lot_match:
                        lor_lot_match = re.search(r'\b([A-Z]\d{3,5})\b', text)
                    lot = lor_lot_match.group(1).strip() if lor_lot_match else None
                    
                    code_match = re.search(r'Item\s*Number[\s\n\|:]*(?!Lot|Code)([A-Za-z0-9-]+)', text, re.IGNORECASE)
                    if not code_match:
                        code_match = re.search(r'\b(\d{3,4})\b', text)
                    lor_mfr_code = code_match.group(1).strip() if code_match else mfr_code
                    
                    lor_flavor_match = re.search(r'Item[\s\n\|:]*([A-Za-z0-9\-\s\(\)\'\&\+,\.%#]{2,60}?)(?=\n|  |$|\s+Item\s+Number|\s+Lot\s+Code)', text, re.IGNORECASE)
                    if lor_flavor_match:
                        raw_lor_flavor = lor_flavor_match.group(1).strip()
                        clean_lor_flavor = re.sub(r'(?i)\s+FLAVOR\b|\s+FLVR\b', '', raw_lor_flavor).strip()
                        lor_clean_flavor = re.sub(r'[\\/*?:"<>|]', "", clean_lor_flavor).strip().replace(" ", "_")
                    else:
                        lor_clean_flavor = universal_clean_flavor
                        
                    lor_mfg_match = re.search(r'(?:Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday),\s+([A-Za-z]+\s+\d{1,2},\s+\d{4})', text, re.IGNORECASE)
                    lor_mfr_date = lor_mfg_match.group(1).strip() if lor_mfg_match else mfr_date
                    lor_bb_match = re.search(r'\b(Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?|Sun)\s+20\d{2}\b', text, re.IGNORECASE)
                    lor_bb_date = lor_bb_match.group(0).strip() if lor_bb_match else bb_date
                    clean_date = lor_mfr_date.replace(", ", "-").replace(" ", "-") if lor_mfr_date != "N/A" else "Unknown_Date"

                    page_data.append({
                        'is_capella': False, 'is_flavor_west': False, 'is_hedessent': False, 'is_flavorah': False, 'is_lorann': True, 'is_purilum': False, 'is_sobucky': False, 'is_tfa': False, 'is_wf': False,
                        'lot': lot, 'flavor': lor_clean_flavor, 'date': clean_date,
                        'mfr_code': lor_mfr_code, 'mfr_date': lor_mfr_date, 'bb_date': lor_bb_date
                    })

                elif is_purilum:
                    pur_lot_match = re.search(r'Batch\s*#:[\s\n\|]*([A-Za-z0-9-]+)', text, re.IGNORECASE)
                    lot = pur_lot_match.group(1).strip() if pur_lot_match else None
                    pur_code_match = re.search(r'Material:[\s\n\|]*([A-Za-z0-9-]+)', text, re.IGNORECASE)
                    pur_mfr_code = pur_code_match.group(1).strip() if pur_code_match else mfr_code
                    pur_mfg_match = re.search(r'Date\s*of\s*Mfg:[\s\n\|]*([0-9/-]+)', text, re.IGNORECASE)
                    pur_mfr_date = pur_mfg_match.group(1).strip() if pur_mfg_match else mfr_date
                    pur_bb_match = re.search(r'Expiry:[\s\n\|]*([0-9/-]+)', text, re.IGNORECASE)
                    pur_bb_date = pur_bb_match.group(1).strip() if pur_bb_match else bb_date
                    clean_date = pur_mfr_date.replace("/", "-") if pur_mfr_date != "N/A" else "Unknown_Date"
                    
                    pur_flavor_match = re.search(r'Product\s*Name:[\s\n\|]*([\s\S]+?)(?=\n\s*Material:)', text, re.IGNORECASE)
                    if pur_flavor_match:
                        raw_pur_flavor = pur_flavor_match.group(1).replace('|', '').strip()
                        raw_pur_flavor = re.sub(r'(?i)\bflavor\b', '', raw_pur_flavor).strip()
                        clean_pur_flavor = re.sub(r'[\\/*?:"<>|]', '', raw_pur_flavor).replace('\n', ' ').strip()
                        pur_clean_flavor = re.sub(r'\s+', '_', clean_pur_flavor)
                    else:
                        pur_clean_flavor = universal_clean_flavor
                        
                    page_data.append({
                        'is_capella': False, 'is_flavor_west': False, 'is_hedessent': False, 'is_flavorah': False, 'is_lorann': False, 'is_purilum': True, 'is_sobucky': False, 'is_tfa': False, 'is_wf': False,
                        'lot': lot, 'flavor': pur_clean_flavor, 'date': clean_date,
                        'mfr_code': pur_mfr_code, 'mfr_date': pur_mfr_date, 'bb_date': pur_bb_date
                    })

                elif is_sobucky:
                    sob_batch_match = re.search(r'BATCH NUMBER:[\s\n\|]*([A-Za-z0-9/-]+)', text, re.IGNORECASE)
                    lot = sob_batch_match.group(1).strip() if sob_batch_match else None
                    sob_code_match = re.search(r'PRODUCT CODE[\s\n\|]*([A-Za-z0-9/-]+)', text, re.IGNORECASE)
                    sob_mfr_code = sob_code_match.group(1).strip() if sob_code_match else mfr_code
                    
                    sob_flavor_match = re.search(r'PRODUCT[\s\n\|]*(?:SUPER AROMAS\s*)?([^\n\(]+)', text, re.IGNORECASE)
                    if sob_flavor_match:
                        raw_sob_flavor = sob_flavor_match.group(1).strip()
                        sob_clean_flavor = re.sub(r'[\\/*?:"<>|]', "", raw_sob_flavor).strip().replace(" ", "_")
                    else:
                        sob_clean_flavor = universal_clean_flavor
                        
                    sob_mfg_match = re.search(r'Date:[\s\n\|]*([0-9]{2}\.[0-9]{2}\.[0-9]{4})', text)
                    sob_mfr_date = sob_mfg_match.group(1).strip() if sob_mfg_match else mfr_date
                    sob_bb_match = re.search(r'BEST BEFORE DATE:[\s\n\|]*([0-9\.]+)', text, re.IGNORECASE)
                    sob_bb_date = sob_bb_match.group(1).strip() if sob_bb_match else bb_date
                    clean_date = sob_mfr_date.replace(".", "-") if sob_mfr_date != "N/A" else "Unknown_Date"

                    page_data.append({
                        'is_capella': False, 'is_flavor_west': False, 'is_hedessent': False, 'is_flavorah': False, 'is_lorann': False, 'is_purilum': False, 'is_sobucky': True, 'is_tfa': False, 'is_wf': False,
                        'lot': lot, 'flavor': sob_clean_flavor, 'date': clean_date,
                        'mfr_code': sob_mfr_code, 'mfr_date': sob_mfr_date, 'bb_date': sob_bb_date
                    })

                elif is_tfa:
                    tfa_mfg_match = re.search(r'(?:Manufacture\s*Date|Date):\s*([0-9]{2}[-/][0-9]{2}[-/][0-9]{4})', text, re.IGNORECASE)
                    tfa_mfr_date = tfa_mfg_match.group(1).strip() if tfa_mfg_match else mfr_date
                    clean_date = tfa_mfr_date.replace("/", "-") if tfa_mfr_date != "N/A" else "Unknown_Date"
                    
                    tfa_lot_match = re.search(r'\b([A-Z]?\d{5,12})\b', text)
                    lot = tfa_lot_match.group(1).strip() if tfa_lot_match else None
                    tfa_mfr_code = lot if lot else mfr_code

                    if "(cid" in text:
                        base_name = os.path.splitext(os.path.basename(pdf_path))[0]
                        clean_base = re.sub(r'(?i)\b(?:tfa|tpa|coa|coas)\b', '', base_name)
                        clean_base = re.sub(r'[\(\)\d\-\_]+$', '', clean_base).strip()
                        tfa_clean_flavor = re.sub(r'[^A-Za-z0-9\s\-_]', '', clean_base).strip().replace(" ", "_")
                    else:
                        if lot:
                            tfa_flavor_match = re.search(r'([\s\S]+?)\n\s*\b' + re.escape(lot) + r'\b', text, re.IGNORECASE)
                            raw_block = tfa_flavor_match.group(1).strip() if tfa_flavor_match else universal_clean_flavor
                        else:
                            raw_block = universal_clean_flavor

                        lines = [line.strip() for line in raw_block.split('\n') if line.strip()]
                        clean_lines = []
                        for line in lines:
                            if re.search(r'APPRENTICE|CERTIFICATE|SPECIFICATION|Contains|Distributor', line, re.IGNORECASE):
                                continue
                            clean_lines.append(line)

                        clean_tfa_flavor = " ".join(clean_lines) if clean_lines else universal_clean_flavor
                        clean_tfa_flavor = re.sub(r'(?i)\s+Flavor(?:\s*\*+)?$', '', clean_tfa_flavor).strip()
                        tfa_clean_flavor = re.sub(r'[^A-Za-z0-9\s\-_]', '', clean_tfa_flavor).strip().replace(" ", "_")

                    tfa_bb_match = re.search(r'BEST\s*IF\s*USED\s*BY\s*([^\n\.]+)', text, re.IGNORECASE)
                    tfa_bb_date = tfa_bb_match.group(1).strip() if tfa_bb_match else "30 Months"

                    page_data.append({
                        'is_capella': False, 'is_flavor_west': False, 'is_hedessent': False, 'is_flavorah': False, 'is_lorann': False, 'is_purilum': False, 'is_sobucky': False, 'is_tfa': True, 'is_wf': False,
                        'lot': lot, 'flavor': tfa_clean_flavor, 'date': clean_date,
                        'mfr_code': tfa_mfr_code, 'mfr_date': tfa_mfr_date, 'bb_date': tfa_bb_date
                    })

                elif is_wf:
                    # --- WONDER FLAVOURS (WF) RULES UPGRADED ---
                    wf_batch_match = re.search(r'Batch\s*Number:[\s\n\|]*([A-Za-z0-9-]+)', text, re.IGNORECASE)
                    lot = wf_batch_match.group(1).strip() if wf_batch_match else None
                    
                    wf_code_match = re.search(r'Product\s*Code:[\s\n\|]*([A-Za-z0-9-]+)', text, re.IGNORECASE)
                    wf_mfr_code = wf_code_match.group(1).strip() if wf_code_match else mfr_code
                    
                    # Target multiline block between 'Product Name:' and the next label
                    wf_flavor_match = re.search(r'Product\s*Name:[\s\n\|]*([\s\S]+?)(?=\n\s*(?:Sample|Product\s*Code|Customer|Batch))', text, re.IGNORECASE)
                    if wf_flavor_match:
                        raw_wf_flavor = wf_flavor_match.group(1).replace('\n', ' ').strip()
                        clean_wf_flavor = re.sub(r'(?i)\s+Flavour\b|\s+Flavor\b', '', raw_wf_flavor).strip()
                        wf_clean_flavor = re.sub(r'[^A-Za-z0-9\s\-_]', '', clean_wf_flavor).strip().replace(" ", "_")
                        wf_clean_flavor = re.sub(r'_+', '_', wf_clean_flavor)
                    else:
                        wf_clean_flavor = universal_clean_flavor
                        
                    wf_mfg_match = re.search(r'Manufacture\s*Date:[\s\n\|]*([0-9/-]+)', text, re.IGNORECASE)
                    wf_mfr_date = wf_mfg_match.group(1).strip() if wf_mfg_match else mfr_date
                    
                    wf_bb_match = re.search(r'Best\s*Before\s*/\s*Retest\s*Date:[\s\n\|]*([0-9/-]+)', text, re.IGNORECASE)
                    wf_bb_date = wf_bb_match.group(1).strip() if wf_bb_match else bb_date
                    
                    clean_date = wf_mfr_date.replace("/", "-") if wf_mfr_date != "N/A" else "Unknown_Date"

                    page_data.append({
                        'is_capella': False, 'is_flavor_west': False, 'is_hedessent': False, 'is_flavorah': False, 'is_lorann': False, 'is_purilum': False, 'is_sobucky': False, 'is_tfa': False, 'is_wf': True,
                        'lot': lot, 'flavor': wf_clean_flavor, 'date': clean_date,
                        'mfr_code': wf_mfr_code, 'mfr_date': wf_mfr_date, 'bb_date': wf_bb_date
                    })

                else:
                    lot_match = re.search(r'(?:Lot|Batch)\s*(?:Code|Number|No|#)?[\s\n\|:]*([A-Za-z0-9-]+)', text, re.IGNORECASE)
                    lot = lot_match.group(1).strip() if lot_match else None
                    
                    page_data.append({
                        'is_capella': False, 'is_flavor_west': False, 'is_hedessent': False, 'is_flavorah': False, 'is_lorann': False, 'is_purilum': False, 'is_sobucky': False, 'is_tfa': False, 'is_wf': False,
                        'lot': lot, 'flavor': universal_clean_flavor, 'date': None,
                        'mfr_code': mfr_code, 'mfr_date': mfr_date, 'bb_date': bb_date
                    })
    except Exception as e:
        print(f"Error reading {pdf_path}: {e}")
        
    return page_data

def generate_filename(page_dict, original_pdf_path=""):
    """Creates the file name based on brand specifications, scrubbing font CID glitches and special characters."""
    if not page_dict['lot']:
        return None
    
    raw_lot = str(page_dict['lot'])
    raw_flavor = str(page_dict['flavor'])
    
    clean_lot = re.sub(r'\(cid\d+\)', '', raw_lot)
    clean_flavor = re.sub(r'\(cid\d+\)', '', raw_flavor)
    
    clean_lot = re.sub(r'[^A-Za-z0-9\s\-_]', '', clean_lot)
    clean_flavor = re.sub(r'[^A-Za-z0-9\s\-_]', '', clean_flavor)
    
    clean_lot = re.sub(r'[\s_]+', '_', clean_lot).strip('_')
    clean_flavor = re.sub(r'[\s_]+', '_', clean_flavor).strip('_')
    
    if not clean_flavor and original_pdf_path:
        base_name = os.path.splitext(os.path.basename(original_pdf_path))[0]
        clean_flavor = re.sub(r'[^A-Za-z0-9\s\-_]', '', base_name)
        clean_flavor = re.sub(r'[\s_]+', '_', clean_flavor).strip('_')
        
    if not clean_flavor:
        clean_flavor = "Extracted_COA"
        
    if (page_dict.get('is_capella') or page_dict.get('is_flavor_west') or 
        page_dict.get('is_hedessent') or page_dict.get('is_flavorah') or 
        page_dict.get('is_lorann') or page_dict.get('is_purilum') or 
        page_dict.get('is_sobucky') or page_dict.get('is_tfa') or page_dict.get('is_wf')):
        return f"{clean_flavor}_{clean_lot}_{page_dict['date']}.pdf"
    else:
        return f"lot_{clean_lot}.pdf"

def process_vault(source_directory, output_directory):
    os.makedirs(output_directory, exist_ok=True)
    review_folder = os.path.join(output_directory, "_Needs_Review")
    os.makedirs(review_folder, exist_ok=True)

    log_records = [] 

    for root, dirs, files in os.walk(source_directory):
        relative_path = os.path.relpath(root, source_directory)
        target_dir = os.path.join(output_directory, relative_path)
        
        if relative_path != ".":
            os.makedirs(target_dir, exist_ok=True)

        current_folder_name = os.path.basename(root)

        for file in files:
            if not file.lower().endswith('.pdf'):
                continue 
                
            original_file_path = os.path.join(root, file)
            print(f"\nProcessing: {original_file_path}")
            
            pages = extract_page_data(original_file_path, current_folder_name)
            valid_pages = [p for p in pages if p['lot'] is not None]
            unique_lots = list({p['lot']: p for p in valid_pages}.values())

            if len(unique_lots) == 0:
                shutil.copy(original_file_path, os.path.join(review_folder, file))
                print(f" -> Flagged for review")
                log_records.append([file, "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "Moved to Review"])
                continue

            if len(unique_lots) == 1:
                page_info = unique_lots[0]
                new_name = generate_filename(page_info, original_file_path)
                
                shutil.copy(original_file_path, os.path.join(target_dir, new_name))
                print(f" -> Copied as single COA: {new_name}")
                
                log_records.append([
                    file, new_name, 
                    page_info['flavor'] or "N/A", 
                    page_info['date'] or "N/A", 
                    page_info['lot'], 
                    page_info['mfr_code'],
                    page_info['mfr_date'],
                    page_info['bb_date'],
                    "Saved Correctly"
                ])
                continue

            reader = PdfReader(original_file_path)
            current_writer = None
            current_page_info = None
            
            for page_idx, page_info in enumerate(pages):
                if page_info['lot'] and (not current_page_info or page_info['lot'] != current_page_info['lot']):
                    if current_writer and current_page_info:
                        new_name = generate_filename(current_page_info, original_file_path)
                        out_path = os.path.join(target_dir, new_name)
                        with open(out_path, "wb") as f:
                            current_writer.write(f)
                        
                        log_records.append([
                            file, new_name, 
                            current_page_info['flavor'] or "N/A", 
                            current_page_info['date'] or "N/A", 
                            current_page_info['lot'], 
                            current_page_info['mfr_code'],
                            current_page_info['mfr_date'],
                            current_page_info['bb_date'],
                            "Split & Saved Correctly"
                        ])
                    
                    current_writer = PdfWriter()
                    current_page_info = page_info
                
                if current_writer:
                    current_writer.add_page(reader.pages[page_idx])

            if current_writer and current_page_info:
                new_name = generate_filename(current_page_info, original_file_path)
                out_path = os.path.join(target_dir, new_name)
                with open(out_path, "wb") as f:
                    current_writer.write(f)
                
                log_records.append([
                    file, new_name, 
                    current_page_info['flavor'] or "N/A", 
                    current_page_info['date'] or "N/A", 
                    current_page_info['lot'], 
                    current_page_info['mfr_code'],
                    current_page_info['mfr_date'],
                    current_page_info['bb_date'],
                    "Split & Saved Correctly"
                ])
                
            print(f" -> Split combined file into {len(unique_lots)} individual COAs.")

    # --- Excel Export Logic ---
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "COA Log"
    
    ws.append([
        "Original File", "New File Name", "Flavor", "Date", "Lot / Lab", 
        "Mfr Code", "Mfr Date", "Best By Date", "Status"
    ])

    header_fill = PatternFill(start_color="2F75B5", end_color="2F75B5", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in log_records:
        ws.append(row)

    excel_path = os.path.join(output_directory, "COA_Processing_Log.xlsx")
    wb.save(excel_path)
    print(f"\nExcel log created: {excel_path}")

# --- DESKTOP PATHING ---
source_folder = "C:/Users/willi/OneDrive/Desktop/Messy_Test_Folder"
output_folder = "C:/Users/willi/OneDrive/Desktop/Cleaned_Vault_Folder"

print("Starting COA extraction...")
process_vault(source_folder, output_folder)
print("\nProcess complete! Check the Cleaned_Vault_Folder on your desktop.")